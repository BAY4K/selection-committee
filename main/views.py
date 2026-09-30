from datetime import datetime
from django.conf import settings
from django.contrib import messages
from django.contrib.admin.views.decorators import staff_member_required
from django.views.decorators.http import require_POST
from django.db.models import Model
import string

import openpyxl
from django.http import HttpResponse, FileResponse, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse_lazy
from django.views.generic import TemplateView, ListView, FormView, CreateView
from openpyxl.workbook import Workbook

from main.admin import StudentResource
from main.forms import ApplicantShortForm, FieldSelectionForm
from main.models import Department, Applicant, Parent, School, ApplicantAdmissionView, Document, Admission
from main.utils import parse_schools, fill_template


# Create your views here.
class IndexView(TemplateView):
    template_name = 'main/index.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['departments'] = Department.objects.all()
        return context


class TestFormView(CreateView):
    template_name = 'main/form.html'
    form_class = ApplicantShortForm
    success_url = reverse_lazy('home')

    def form_valid(self, form):
        response = super().form_valid(form)
        messages.success(self.request, 'Заявка отправлена! После рассмотрения мы пришлём данные для входа на указанную почту.')
        return response


@staff_member_required
def download_table(request):
    selection = {group + '_CHOICES': [name for name, _ in getattr(FieldSelectionForm, group + '_FIELDS')]
                 for group in ('APPLICANT', 'DOCUMENT', 'PARENT', 'ADMISSION')}
    return export_to_excel(request, selection=selection)


def download_document(request):
    # Define text file name
    filename = 'tipovaya-forma_soglasie.doc'
    # Define the full file path
    filepath = settings.BASE_DIR / 'static' / 'documents' / filename
    # Open the file for reading content
    response = FileResponse(open(filepath, 'rb'))
    response['Content-Disposition'] = 'attachment; filename="your_file.doc"'
    # Return the response value
    return response


@staff_member_required
@require_POST
def update_schools(request):
    parse_schools()
    return redirect('home')


@staff_member_required
def generate_document(request, person_id):
    person = get_object_or_404(Applicant, id=person_id)
    template_path = settings.BASE_DIR / 'static' / 'documents' / 'ZayavlenieAbiturienta.docx'
    filled_doc = fill_template(person.id, template_path)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    response['Content-Disposition'] = f'attachment; filename="document.docx"'
    filled_doc.save(response)

    return response


def normalize(text):
    return ''.join(char for char in text if char not in string.punctuation).lower()


def autocomplete(request):
    term = request.GET.get('term', '').strip()
    names = School.objects.filter(name__icontains=term).order_by('name').values_list('name', flat=True)[:30] if term else []
    return JsonResponse(list(names), safe=False)


@staff_member_required
def export_data(request):
    return export_to_excel(request)


@staff_member_required
def export_students(request):
    dataset = StudentResource().export()
    for index, row in enumerate(dataset):
        dataset[index] = ["'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value for value in row]
    response = HttpResponse(dataset.xlsx,
                            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="students.xlsx"'
    return response


@staff_member_required
def export_to_excel(request, selection=None):
    if request.method == 'POST' or selection is not None:
        form = FieldSelectionForm(selection if selection is not None else request.POST)
        if form.is_valid():
            selected_fields = {
                'applicant': form.cleaned_data.get('APPLICANT_CHOICES', []),
                'document': form.cleaned_data.get('DOCUMENT_CHOICES', []),
                'parent': form.cleaned_data.get('PARENT_CHOICES', []),
                'admission': form.cleaned_data.get('ADMISSION_CHOICES', [])
            }

            # Create a workbook and add a worksheet
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Applicant Data"

            # Add headers
            row_num = 1
            col_num = 1
            headers = [('ID', 'ID')]  # Add 'ID' as the first column
            verbose_names = ['ID']
            order = ['last_name', 'first_name', 'patronymic', 'gender', 'birth_date', 'email', 'phone', 'address',
                     'school', 'graduation_date', 'education', 'consent', 'SNILS',
                     'INN', 'passport_number', 'mother_full_name',
                     'mother_phone', 'father_full_name', 'father_phone', 'admission_date', 'number_of_5', 'number_of_4',
                     'number_of_3', 'average_score', 'internal_exam', 'application_status', 'original_or_copy',
                     'out_of_budget', 'received_receipt', 'internal_exam_conducted', 'documents_collected',
                     'application_in_gov_services']

            def sort_fields(fields):
                return sorted(fields, key=lambda x: order.index(x[0]) if x[0] in order else len(order))

            # Map selected fields to their verbose names
            models_map = {
                'applicant': Applicant,
                'document': Document,
                'parent': Parent,
                'admission': Admission,
            }

            for model, fields in selected_fields.items():
                sorted_fields = sort_fields([(field, field) for field in fields])
                for field, _ in sorted_fields:
                    headers.append((model, field))
                    verbose_name = models_map[model]._meta.get_field(field).verbose_name
                    verbose_names.append(verbose_name)
                    ws.cell(row=row_num, column=col_num + 1, value=verbose_name)  # Shift by 1 for ID
                    col_num += 1

            # Add headers for ID
            ws.cell(row=row_num, column=1, value='ID')

            # Add data
            for idx, obj in enumerate(Applicant.objects.select_related('document', 'parents', 'student').all(), start=1):
                row_num += 1
                col_num = 1
                # Add ID
                ws.cell(row=row_num, column=col_num, value=idx)
                col_num += 1
                for model, field in headers[1:]:  # Skip the first ID header
                    related_obj = obj if model == 'applicant' else getattr(obj, {'parent': 'parents', 'admission': 'student'}.get(model, model), None)
                    value = getattr(related_obj, field, None)
                    if callable(value):
                        value = value()
                    if isinstance(value, Model):
                        value = str(value)
                    if isinstance(value, datetime) and value.tzinfo is not None:
                        from django.utils import timezone
                        value = timezone.make_naive(value, timezone.get_current_timezone())
                    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
                        value = "'" + value
                    if isinstance(value, bool):
                        value = 'Да' if value else 'Нет'
                    ws.cell(row=row_num, column=col_num, value=value)
                    col_num += 1

            # Adjust column widths
            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter  # Get the column name
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = (max_length + 2)
                ws.column_dimensions[column].width = adjusted_width

            # Adjust row heights
            for row in ws.iter_rows():
                max_height = 15  # Default height
                for cell in row:
                    if cell.value:
                        cell_height = (len(str(cell.value)) // 20) * 5
                        if cell_height > max_height:
                            max_height = cell_height
                ws.row_dimensions[cell.row].height = max_height

            # Save the workbook
            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = 'attachment; filename=applicant_data.xlsx'
            wb.save(response)
            return response
    else:
        form = FieldSelectionForm()

    return render(request, 'admin/export_data.html', {'form': form})


def page_not_found(request, exception):
    return HttpResponse("404 NOT FOUND", status=404)
