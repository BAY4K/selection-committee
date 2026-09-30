from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.test import TestCase, override_settings
from django.urls import reverse, resolve
from openpyxl import load_workbook
from docx import Document as WordDocument

from account.models import User
from main.models import Applicant, Admission, Department, Document, Parent, School
from main.utils import create_account, deny_student, warn_student, parse_schools, replace_text_in_runs


def applicant(email='student@example.com', **kwargs):
    return Applicant.objects.create(last_name='Иванов', first_name='Иван', patronymic='Иванович',
                                    birth_date=date(2005, 1, 1), email=email, school='Школа', **kwargs)


class AccessTests(TestCase):
    def setUp(self):
        self.person = applicant()
        self.other = applicant('other@example.com')
        self.user = User.objects.create_user('student', password='test-password', student=self.person)
        for person in (self.person, self.other):
            Admission.objects.create(applicant=person)
            Document.objects.create(student=person)
            Parent.objects.create(student=person)

    def test_anonymous_profile_redirects(self):
        for route in ('profile-edit', 'parents-edit', 'documents-edit', 'admission-create'):
            self.assertEqual(self.client.get(reverse('account:' + route, args=[self.person.pk])).status_code, 302)

    def test_foreign_profiles_are_inaccessible(self):
        self.client.force_login(self.user)
        for route in ('profile-edit', 'parents-edit', 'documents-edit', 'admission-create'):
            url = reverse('account:' + route, args=[self.other.pk])
            self.assertEqual(self.client.get(url).status_code, 404)
            self.assertEqual(self.client.post(url, {'phone': 'changed'}).status_code, 404)

    def test_own_profiles_render(self):
        self.client.force_login(self.user)
        for route in ('profile-edit', 'parents-edit', 'documents-edit', 'admission-create'):
            self.assertEqual(self.client.get(reverse('account:' + route, args=[self.person.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse('account:profile', args=[self.user.pk])).status_code, 200)

    def test_exports_require_staff(self):
        self.client.force_login(self.user)
        for url in ('/download-table/', '/export/', '/admin/export/', '/admin/export_students/', '/generate_document/1/', '/update-schools/'):
            self.assertEqual(self.client.get(url).status_code, 302)

    def test_school_update_requires_post(self):
        self.user.is_staff = True
        self.user.save()
        self.client.force_login(self.user)
        self.assertEqual(self.client.get('/update-schools/').status_code, 405)
        with patch('main.views.parse_schools') as parse:
            self.assertEqual(self.client.post('/update-schools/').status_code, 302)
            parse.assert_called_once()

    def test_no_applicant_profile_returns_404(self):
        user = User.objects.create_user('no-profile')
        self.client.force_login(user)
        self.assertEqual(self.client.get(reverse('account:profile', args=[user.pk])).status_code, 404)

    def test_inactive_email_login_rejected(self):
        from account.authentication import EmailAuthBackend
        self.user.email = 'login@example.com'
        self.user.is_active = False
        self.user.save()
        backend = EmailAuthBackend()
        self.assertIsNone(backend.authenticate(None, username=self.user.email, password='test-password'))
        self.assertIsNone(backend.get_user(self.user.pk))

    def test_password_reset_pages_render(self):
        for route in ('password_reset', 'password_reset_done', 'password_reset_complete'):
            self.assertEqual(self.client.get(reverse('account:' + route)).status_code, 200)

    def test_export_students_route_is_reachable(self):
        self.assertEqual(resolve('/admin/export_students/').url_name, 'export_students')

    def test_password_change_flow(self):
        self.client.force_login(self.user)
        url = reverse('account:password_change')
        self.assertEqual(self.client.get(url).status_code, 200)
        response = self.client.post(url, {'old_password': 'test-password', 'new_password1': 'NewPassWord!9371', 'new_password2': 'NewPassWord!9371'})
        self.assertRedirects(response, reverse('account:password_change_done'))
        self.user.refresh_from_db()
        self.assertTrue(self.user.check_password('NewPassWord!9371'))

    def test_staff_student_export(self):
        self.user.is_staff = True
        self.user.save()
        self.client.force_login(self.user)
        response = self.client.get('/admin/export_students/')
        self.assertEqual(response.status_code, 200)
        self.assertGreater(len(load_workbook(BytesIO(response.content)).active['A']), 1)


class DataTests(TestCase):
    def test_empty_and_cleared_grades(self):
        admission = Admission.objects.create(applicant=applicant(), number_of_5=None, number_of_4=None, number_of_3=None, internal_exam=None)
        self.assertEqual(admission.average_score, Decimal(0))
        self.assertFalse(admission.internal_exam_conducted)
        admission.number_of_5 = 2
        admission.number_of_4 = 1
        admission.save()
        admission.refresh_from_db()
        self.assertEqual(admission.average_score, Decimal('4.67'))
        admission.number_of_5 = admission.number_of_4 = 0
        admission.save(update_fields=['number_of_5', 'number_of_4'])
        admission.refresh_from_db()
        self.assertEqual(admission.average_score, 0)

    def test_deny_and_warn_have_distinct_statuses(self):
        admission = Admission.objects.create(applicant=applicant())
        deny_student(admission)
        self.assertEqual(admission.application_status, 'Отказано')
        warn_student(admission)
        self.assertEqual(admission.application_status, 'Отправлен на заполнение')
        self.assertNotEqual(admission.get_status(), 'None')

    @override_settings(EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
    def test_account_creation_is_complete_and_idempotent(self):
        person = applicant()
        self.assertTrue(create_account(person))
        self.assertTrue(User.objects.get(student=person).has_usable_password())
        self.assertTrue(Admission.objects.filter(applicant=person).exists())
        self.assertTrue(Document.objects.filter(student=person).exists())
        self.assertTrue(Parent.objects.filter(student=person).exists())
        self.assertFalse(create_account(person))

    def test_email_failure_rolls_back_account_creation(self):
        person = applicant()
        with patch('main.utils.send_mail', side_effect=RuntimeError('SMTP failure')):
            with self.assertRaises(RuntimeError):
                create_account(person)
        self.assertFalse(User.objects.filter(student=person).exists())
        self.assertFalse(Admission.objects.filter(applicant=person).exists())
        person.refresh_from_db()
        self.assertEqual(person.status, 'Рассмотрение')

    def test_failed_school_refresh_preserves_data(self):
        School.objects.create(name='Existing')
        with patch('main.utils.parse_info', side_effect=RuntimeError('Network failure')):
            with self.assertRaises(RuntimeError):
                parse_schools()
        self.assertTrue(School.objects.filter(name='Existing').exists())

    def test_autocomplete_without_term_and_literal_punctuation(self):
        School.objects.create(name='Школа № 1 (центр)')
        self.assertEqual(self.client.get('/autocomplete/').json(), [])
        self.assertEqual(self.client.get('/autocomplete/', {'term': '(центр)'}).json(), ['Школа № 1 (центр)'])

    def test_export_handles_incomplete_applicants_and_selected_groups(self):
        person = applicant()
        person.last_name = '=1+1'
        person.save()
        staff = User.objects.create_user('staff', is_staff=True)
        self.client.force_login(staff)
        response = self.client.get('/download-table/')
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content))
        self.assertEqual(book.active.cell(2, 2).value, "'=1+1")
        self.assertEqual(self.client.post('/export/', {'APPLICANT_CHOICES': ['last_name']}).status_code, 200)

    def test_split_word_placeholders_and_missing_values(self):
        doc = WordDocument()
        paragraph = doc.add_paragraph()
        paragraph.add_run('{last_')
        paragraph.add_run('name} {phone}')
        replace_text_in_runs(paragraph, {'{last_name}': 'Иванов', '{phone}': None})
        self.assertEqual(paragraph.text, 'Иванов ')

    def test_document_generation_handles_missing_documents(self):
        import tempfile
        from pathlib import Path
        from main.utils import fill_template
        person = applicant()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'template.docx'
            doc = WordDocument()
            doc.add_paragraph('{last_name} {first_name} {passport} {issue_date}')
            table = doc.add_table(rows=1, cols=1)
            paragraph = table.cell(0, 0).paragraphs[0]
            paragraph.add_run('{last_')
            paragraph.add_run('name}')
            doc.save(path)
            result = fill_template(person.pk, path)
        self.assertTrue(result.paragraphs[0].text.startswith('Иванов Иван'))
        self.assertEqual(result.tables[0].cell(0, 0).text, 'Иванов')

    def test_passport_parts_match_input_mask(self):
        document = Document(passport_number='12 34 567890')
        self.assertEqual(document.get_passport_series(), '1234')
        self.assertEqual(document.get_passport_num(), '567890')
        document.passport_number = None
        self.assertEqual(document.get_passport_num(), '')

    def test_empty_interviews_do_not_hide_available_students(self):
        from django.utils import timezone
        from main.models import Interview, InternalExam
        from account.forms import InterviewAdminForm, InternalExamAdminForm
        person = applicant()
        interview = Interview.objects.create(interview_date=timezone.now())
        exam = InternalExam.objects.create(exam_date=timezone.now())
        self.assertIn(person, InterviewAdminForm().fields['students'].queryset)
        self.assertIn(person, InternalExamAdminForm().fields['students'].queryset)
        interview.students.add(person)
        exam.students.add(person)
        self.assertIn(person, InterviewAdminForm(instance=interview).fields['students'].queryset)
        self.assertIn(person, InternalExamAdminForm(instance=exam).fields['students'].queryset)
        self.assertNotIn(person, InterviewAdminForm().fields['students'].queryset)

    def test_admission_date_does_not_change_on_edit(self):
        admission = Admission.objects.create(applicant=applicant())
        Admission.objects.filter(pk=admission.pk).update(admission_date=date(2020, 1, 1))
        admission.refresh_from_db()
        admission.number_of_5 = 1
        admission.save()
        admission.refresh_from_db()
        self.assertEqual(admission.admission_date, date(2020, 1, 1))

    def test_legacy_status_migration(self):
        import importlib
        from django.apps import apps
        from django.db import connection
        person = applicant(status='answered', gender='male')
        admission = Admission.objects.create(applicant=person, application_status='denied')
        from types import SimpleNamespace
        importlib.import_module('main.migrations.0008_normalize_legacy_statuses').normalize_statuses(apps, SimpleNamespace(connection=connection))
        person.refresh_from_db()
        admission.refresh_from_db()
        self.assertEqual(person.status, 'Выдан ответ')
        self.assertEqual(person.gender, 'Мужской')
        self.assertEqual(admission.application_status, 'Отказано')

    @override_settings(DEBUG=False)
    def test_missing_page_has_404_status(self):
        self.assertEqual(self.client.get('/does-not-exist/').status_code, 404)

    def test_rank_beyond_first_25_and_russian_status(self):
        department = Department.objects.create(name='Department')
        for index in range(26):
            person = applicant(f'{index}@example.com')
            admission = Admission.objects.create(applicant=person, number_of_4=1, application_status='Принят')
            admission.department.add(department)
        user = User.objects.create_user('rank', student=person)
        self.client.force_login(user)
        response = self.client.get(reverse('account:rank-profile'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['department_admissions'][0][2], 26)
        self.assertFalse(response.context['department_admissions'][0][3])
